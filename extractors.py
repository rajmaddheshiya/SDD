
"""Extract text from uploaded files without saving them to disk."""

import io
import os


TEXT_EXT = {
    ".txt", ".csv", ".tsv", ".log", ".json", ".md", ".xml",
    ".html", ".htm", ".yml", ".yaml", ".env", ".ini", ".cfg",
    ".conf", ".sql", ".py", ".js", ".properties",
}

ALLOWED = TEXT_EXT | {".pdf", ".docx", ".xlsx"}


def get_ext(filename: str) -> str:
    """Return the lowercase extension of a filename."""
    name = os.path.basename(filename).lower()
    extension = os.path.splitext(name)[1]

    if extension:
        return extension

    return name if name.startswith(".") else ""


def decode_text(data: bytes) -> str:
    """Decode text using common encodings."""
    encodings = ("utf-8-sig", "utf-16", "latin-1")

    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue

    return data.decode("utf-8", errors="replace")


def extract(filename: str, data: bytes):
    """Return extracted text units and any processing warnings."""
    extension = get_ext(filename)
    units = []
    warnings = []

    if extension in TEXT_EXT:
        text = decode_text(data)

        for line_number, line in enumerate(text.splitlines(), start=1):
            units.append((f"Line {line_number}", line))

    elif extension == ".pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""

                for line_number, line in enumerate(
                    page_text.splitlines(), start=1
                ):
                    location = f"Page {page_number}, line {line_number}"
                    units.append((location, line))

        if not units:
            warnings.append(
                "No text was found in this PDF. It may contain scanned "
                "images, which require OCR to extract."
            )

    elif extension == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))

        for paragraph_number, paragraph in enumerate(
            document.paragraphs, start=1
        ):
            units.append(
                (f"Paragraph {paragraph_number}", paragraph.text)
            )

        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    location = (
                        f"Table {table_number}, row {row_number}, "
                        f"col {column_number}"
                    )
                    units.append((location, cell.text))

    elif extension == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )

        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            location = (
                                f"{worksheet.title}!{cell.coordinate}"
                            )
                            units.append((location, str(cell.value)))
        finally:
            workbook.close()

    else:
        raise ValueError(
            f"Unsupported file type: {extension or 'unknown'}."
        )

    return units, warnings
